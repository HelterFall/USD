"""Получение курсов USD/CNY и расчёт стоимости грамма золота в RUB."""
from __future__ import annotations

import logging
import os
import xml.etree.ElementTree as ET
from datetime import datetime
from decimal import Decimal
from zoneinfo import ZoneInfo

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CBR_URL = "https://www.cbr.ru/scripts/XML_daily.asp"
GOLD_URL = "https://api.gold-api.com/price/XAU"
TROY_OUNCE_GRAMS = Decimal("31.1034768")
OUTPUT_FILE = os.getenv("OUTPUT_FILE", "exchange_rates.xlsx")
TIMEZONE = os.getenv("TIMEZONE", "Europe/Moscow")
HEADERS = [
    "Время проверки",
    "Дата котировки ЦБ",
    "USD/RUB",
    "CNY/RUB",
    "Золото, USD за тройскую унцию",
    "Золото, RUB за грамм",
    "Источник золота: время обновления",
]

logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")


def get_cbr_rates() -> tuple[datetime, Decimal, Decimal]:
    """Возвращает дату публикации ЦБ и курсы USD/CNY в рублях."""
    response = requests.get(CBR_URL, timeout=20)
    response.raise_for_status()
    root = ET.fromstring(response.content)
    quote_date = datetime.strptime(root.attrib["Date"], "%d.%m.%Y").date()

    rates: dict[str, Decimal] = {}
    for item in root.findall("Valute"):
        code = item.findtext("CharCode")
        if code in {"USD", "CNY"}:
            nominal = Decimal(item.findtext("Nominal", "1").replace(",", "."))
            value = Decimal(item.findtext("Value", "0").replace(",", "."))
            rates[code] = value / nominal

    missing = {"USD", "CNY"} - rates.keys()
    if missing:
        raise ValueError(f"В ответе ЦБ отсутствуют валюты: {', '.join(sorted(missing))}")
    return datetime.combine(quote_date, datetime.min.time()), rates["USD"], rates["CNY"]


def get_gold_price_usd_ounce() -> tuple[Decimal, str]:
    """Получает спотовую котировку золота XAU в USD за тройскую унцию."""
    response = requests.get(GOLD_URL, timeout=20)
    response.raise_for_status()
    payload = response.json()
    price = Decimal(str(payload["price"]))
    updated = str(payload.get("updatedAtReadable") or payload.get("updatedAt") or "не указано")
    if price <= 0:
        raise ValueError("API вернуло некорректную цену золота")
    return price, updated


def save_to_excel(row: list[object], path: str = OUTPUT_FILE) -> None:
    """Добавляет запись в историю, создавая аккуратно оформленный XLSX при необходимости."""
    if os.path.exists(path):
        workbook = load_workbook(path)
        sheet = workbook.active
        if sheet.max_row == 1 and sheet["A1"].value != HEADERS[0]:
            sheet.delete_rows(1, sheet.max_row)
            sheet.append(HEADERS)
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Курсы и золото"
        sheet.append(HEADERS)

    sheet.append(row)
    header_fill = PatternFill("solid", fgColor="17365D")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for cell in sheet[sheet.max_row]:
        cell.alignment = Alignment(vertical="center")
    for cell in (sheet.cell(sheet.max_row, 3), sheet.cell(sheet.max_row, 4),
                 sheet.cell(sheet.max_row, 5), sheet.cell(sheet.max_row, 6)):
        cell.number_format = '#,##0.0000'
    for cell in sheet.iter_rows(min_row=2, min_col=1, max_col=1):
        cell[0].number_format = "yyyy-mm-dd hh:mm:ss"
    for cell in sheet.iter_rows(min_row=2, min_col=2, max_col=2):
        cell[0].number_format = "yyyy-mm-dd"
    widths = [23, 18, 14, 14, 30, 23, 36]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 34
    workbook.save(path)


def main() -> None:
    checked_at = datetime.now(ZoneInfo(TIMEZONE)).replace(tzinfo=None)
    quote_date, usd_rub, cny_rub = get_cbr_rates()
    gold_usd_ounce, gold_updated = get_gold_price_usd_ounce()
    gold_rub_gram = gold_usd_ounce * usd_rub / TROY_OUNCE_GRAMS

    row = [
        checked_at,
        quote_date.date(),
        float(usd_rub),
        float(cny_rub),
        float(gold_usd_ounce),
        float(gold_rub_gram),
        gold_updated,
    ]
    save_to_excel(row)
    logging.info("Данные сохранены в %s", OUTPUT_FILE)
    logging.info("USD/RUB: %s | CNY/RUB: %s", usd_rub, cny_rub)
    logging.info("Золото: %.2f RUB/г (расчётная спотовая цена)", gold_rub_gram)


if __name__ == "__main__":
    try:
        main()
    except (requests.RequestException, ET.ParseError, KeyError, ValueError, OSError) as exc:
        logging.error("Не удалось получить или сохранить данные: %s", exc)
        raise SystemExit(1)
